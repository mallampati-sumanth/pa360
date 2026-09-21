import json
import re
import uuid
from datetime import date, datetime
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from domain_applications.authorizations.models import (
    AuthorizationRequest,
    AuthorizationStatusHistory,
    MissingInformationItem,
)
from domain_applications.documents.models import AuthorizationDocument
from domain_applications.payer_responses.models import PayerResponse
from domain_applications.payers.models import Payer
from domain_applications.patients.models import Patient
from domain_applications.plans.models import Plan
from domain_applications.providers.models import Provider
from domain_applications.services.models import Service


STATUS_MAP = {
    'DRAFT': AuthorizationRequest.Status.DRAFT,
    'SUBMITTED': AuthorizationRequest.Status.SUBMITTED,
    'PENDING': AuthorizationRequest.Status.PENDING,
    'PENDING_REVIEW': AuthorizationRequest.Status.PENDING,
    'APPROVED': AuthorizationRequest.Status.APPROVED,
    'DENIED': AuthorizationRequest.Status.DENIED,
    'ADDITIONAL INFO REQUESTED': AuthorizationRequest.Status.ADDITIONAL_INFO_REQUESTED,
    'PENDING INFO': AuthorizationRequest.Status.ADDITIONAL_INFO_REQUESTED,
}
PRIORITY_MAP = {
    'ROUTINE': AuthorizationRequest.Priority.ROUTINE,
    'NORMAL': AuthorizationRequest.Priority.ROUTINE,
    'URGENT': AuthorizationRequest.Priority.URGENT,
    'EMERGENCY': AuthorizationRequest.Priority.EMERGENCY,
}
RESPONSE_MAP = {
    'APPROVED': PayerResponse.ResponseType.APPROVED,
    'DENIED': PayerResponse.ResponseType.DENIED,
    'PENDING': PayerResponse.ResponseType.PENDING,
    'MORE_INFO_NEEDED': PayerResponse.ResponseType.ADDITIONAL_INFO_REQUESTED,
    'ADDITIONAL INFO REQUESTED': PayerResponse.ResponseType.ADDITIONAL_INFO_REQUESTED,
}
DOCUMENT_MAP = {
    'CLINICAL_NOTE': AuthorizationDocument.DocumentType.CLINICAL_NOTES,
    'CLINICAL_NOTES': AuthorizationDocument.DocumentType.CLINICAL_NOTES,
    'IMAGING': AuthorizationDocument.DocumentType.IMAGING,
    'LAB_RESULT': AuthorizationDocument.DocumentType.LAB_RESULTS,
    'LAB_RESULTS': AuthorizationDocument.DocumentType.LAB_RESULTS,
    'PRESCRIPTION': AuthorizationDocument.DocumentType.OTHER,
}


def clean(value):
    if value is None:
        return ''
    value = str(value).strip()
    if value.upper() in {'', 'N/A', 'NULL', 'NONE', 'NAN'}:
        return ''
    return re.sub(r'\s+', ' ', value)


def key(value, fallback):
    value = clean(value).replace(' ', '')
    return value or fallback


def parse_date(value, default=None):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    value = clean(value)
    if not value:
        return default
    for fmt in ('%Y-%m-%d', '%m/%d/%Y', '%m-%d-%Y'):
        try:
            return datetime.strptime(value, fmt).date()
        except ValueError:
            continue
    return default


def authorization_required(value):
    value = clean(value).upper()
    if value in {'Y', 'YES', 'TRUE', '1'}:
        return True
    if value in {'N', 'NO', 'FALSE', '0'}:
        return False
    return True


def rows(workbook, sheet_name):
    sheet = workbook[sheet_name]
    values = list(sheet.iter_rows(values_only=True))
    headers = [clean(value) for value in values[0]]
    return [dict(zip(headers, row)) for row in values[1:] if any(value is not None for value in row)]


class Command(BaseCommand):
    help = 'Load the attached PA360 workbook into the local SQLite demo database.'

    def add_arguments(self, parser):
        parser.add_argument('workbook', type=Path)
        parser.add_argument('--replace', action='store_true', help='Delete existing demo records before loading.')

    @transaction.atomic
    def handle(self, *args, **options):
        workbook_path = options['workbook']
        if not workbook_path.exists():
            raise CommandError(f'Workbook not found: {workbook_path}')

        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        required_sheets = {
            'Patients', 'Providers', 'Payers', 'Plans', 'Services',
            'Authorization_Requests', 'Sample_Clinical_Documents', 'Payer_Responses',
        }
        missing = required_sheets - set(workbook.sheetnames)
        if missing:
            raise CommandError(f'Missing workbook sheets: {sorted(missing)}')

        if options['replace']:
            PayerResponse.objects.all().delete()
            AuthorizationDocument.objects.all().delete()
            AuthorizationRequest.objects.all().delete()
            Patient.objects.all().delete()
            Plan.objects.all().delete()
            Service.objects.all().delete()
            Provider.objects.all().delete()
            Payer.objects.all().delete()

        payers = {}
        for index, row in enumerate(rows(workbook, 'Payers'), start=1):
            payer_id = key(row.get('Payer_ID'), f'PYR_DEMO_{index:03d}')
            payer, _ = Payer.objects.update_or_create(
                payer_id=payer_id,
                defaults={
                    'payer_name': clean(row.get('Payer_Name')) or f'Demo Payer {index}',
                    'contact_info': {'contact': clean(row.get('Contact_Info'))},
                    'is_active': True,
                },
            )
            payers[payer_id] = payer
        if not payers:
            payer = Payer.objects.create(payer_id='PYR_DEMO_001', payer_name='Demo Payer', contact_info={})
            payers[payer.payer_id] = payer

        plans = {}
        for index, row in enumerate(rows(workbook, 'Plans'), start=1):
            plan_id = key(row.get('Plan_ID'), f'PLN_DEMO_{index:03d}')
            payer_id = key(row.get('Payer_ID'), next(iter(payers)))
            payer = payers.get(payer_id) or payers[next(iter(payers))]
            plan, _ = Plan.objects.update_or_create(
                plan_id=plan_id,
                defaults={
                    'plan_name': clean(row.get('Plan_Name')) or f'Demo Plan {index}',
                    'payer': payer,
                    'pa_rules_summary': {'source_value': clean(row.get('PA_Rules_Summary'))},
                    'copay_info': {},
                    'effective_date': date(2025, 1, 1),
                    'is_active': True,
                },
            )
            plans[plan_id] = plan
        if not plans:
            plan = Plan.objects.create(plan_id='PLN_DEMO_001', plan_name='Demo Plan', payer=next(iter(payers.values())), effective_date=date(2025, 1, 1))
            plans[plan.plan_id] = plan

        providers = {}
        for index, row in enumerate(rows(workbook, 'Providers'), start=1):
            provider_id = key(row.get('Provider_ID'), f'PRV_DEMO_{index:03d}')
            npi = key(row.get('NPI_Number'), f'999999{index:04d}')[:20]
            if Provider.objects.exclude(provider_id=provider_id).filter(npi_number=npi).exists():
                npi = f'{npi[:16]}{index:04d}'[:20]
            provider, _ = Provider.objects.update_or_create(
                provider_id=provider_id,
                defaults={
                    'provider_name': clean(row.get('Provider_Name')) or f'Demo Provider {index}',
                    'specialty': clean(row.get('Specialty')) or 'General Practice',
                    'npi_number': npi,
                    'practice_location': clean(row.get('Practice_Location')) or 'Demo Location',
                    'is_active': True,
                },
            )
            providers[provider_id] = provider

        services = {}
        for index, row in enumerate(rows(workbook, 'Services'), start=1):
            service_code = key(row.get('Service_Code'), f'SVC_DEMO_{index:03d}')
            service, _ = Service.objects.update_or_create(
                service_code=service_code,
                defaults={
                    'service_name': clean(row.get('Service_Name')) or f'Demo Service {index}',
                    'specialty': clean(row.get('Specialty')) or 'General',
                    'authorization_required': authorization_required(row.get('Authorization_Required')),
                    'is_active': True,
                },
            )
            services[service_code] = service

        patients = {}
        for index, row in enumerate(rows(workbook, 'Patients'), start=1):
            patient_id = key(row.get('Patient_ID'), f'PAT_DEMO_{index:04d}')
            plan_id = key(row.get('Plan_ID'), next(iter(plans)))
            plan = plans.get(plan_id) or plans[next(iter(plans))]
            gender = clean(row.get('Gender')).upper()[:1]
            patient, _ = Patient.objects.update_or_create(
                patient_id=patient_id,
                defaults={
                    'name': clean(row.get('Name')) or f'Demo Patient {index}',
                    'date_of_birth': parse_date(row.get('Date_of_Birth'), date(1980, 1, 1)),
                    'gender': gender if gender in {'M', 'F', 'O'} else 'O',
                    'insurance_id': clean(row.get('Insurance_ID')),
                    'plan': plan,
                    'contact_number': clean(row.get('Contact_Number')),
                    'address': clean(row.get('Address')) or 'Demo Address',
                    'enrollment_status': Patient.EnrollmentStatus.ACTIVE,
                },
            )
            patients[patient_id] = patient

        authorizations = {}
        for index, row in enumerate(rows(workbook, 'Authorization_Requests'), start=1):
            authorization_key = key(row.get('Authorization_ID'), f'AUTH_DEMO_{index:05d}')
            authorization_id = uuid.uuid5(uuid.NAMESPACE_URL, f'pa360-demo:{authorization_key}')
            patient_id = key(row.get('Patient_ID'), next(iter(patients)))
            provider_id = key(row.get('Provider_ID'), next(iter(providers)))
            payer_id = key(row.get('Payer_ID'), next(iter(payers)))
            plan_id = key(row.get('Plan_ID'), next(iter(plans)))
            service_code = key(row.get('Service_Code'), next(iter(services)))
            authorization, _ = AuthorizationRequest.objects.update_or_create(
                authorization_id=authorization_id,
                defaults={
                    'patient': patients.get(patient_id) or next(iter(patients.values())),
                    'provider': providers.get(provider_id) or next(iter(providers.values())),
                    'payer': payers.get(payer_id) or next(iter(payers.values())),
                    'plan': plans.get(plan_id) or next(iter(plans.values())),
                    'service_code': services.get(service_code) or next(iter(services.values())),
                    'diagnosis': clean(row.get('Diagnosis')) or 'Demo diagnosis',
                    'clinical_indication': clean(row.get('Clinical_Indication')) or 'Demo clinical indication',
                    'request_date': parse_date(row.get('Request_Date'), date.today()),
                    'status': STATUS_MAP.get(clean(row.get('Status')).upper(), AuthorizationRequest.Status.PENDING),
                    'priority': PRIORITY_MAP.get(clean(row.get('Priority')).upper(), AuthorizationRequest.Priority.ROUTINE),
                    'missing_information': [clean(row.get('Missing_Information'))] if clean(row.get('Missing_Information')) else [],
                    'supporting_document': clean(row.get('Supporting_Document')),
                    'submission_date': parse_date(row.get('Submission_Date')),
                    'decision_date': parse_date(row.get('Decision_Date')),
                },
            )
            AuthorizationStatusHistory.objects.get_or_create(
                authorization_request=authorization,
                status=authorization.status,
            )
            missing_value = clean(row.get('Missing_Information'))
            for item_name in [item.strip() for item in missing_value.split(',') if item.strip()]:
                MissingInformationItem.objects.update_or_create(
                    authorization_request=authorization,
                    item_name=item_name,
                    defaults={'status': MissingInformationItem.ItemStatus.MISSING},
                )
            authorizations[authorization_key] = authorization

        for index, row in enumerate(rows(workbook, 'Sample_Clinical_Documents'), start=1):
            authorization_key = key(row.get('Authorization_ID'), '')
            authorization = authorizations.get(authorization_key)
            if not authorization:
                continue
            document_key = key(row.get('Document_ID'), f'DOC_DEMO_{index:04d}')
            document_id = uuid.uuid5(uuid.NAMESPACE_URL, f'pa360-document:{document_key}')
            document_type = DOCUMENT_MAP.get(clean(row.get('Document_Type')).upper(), AuthorizationDocument.DocumentType.OTHER)
            AuthorizationDocument.objects.update_or_create(
                document_id=document_id,
                defaults={
                    'authorization_request': authorization,
                    'document_type': document_type,
                    'file_name': f'{document_key}.txt',
                    'file_path': f'demo://clinical-documents/{document_key}.txt',
                    'file_size': len(clean(row.get('Document_Text')).encode('utf-8')),
                    'mime_type': 'text/plain',
                    'extracted_text': clean(row.get('Document_Text')),
                    'extraction_status': AuthorizationDocument.ExtractionStatus.COMPLETED,
                },
            )

        for index, row in enumerate(rows(workbook, 'Payer_Responses'), start=1):
            authorization_key = key(row.get('Authorization_ID'), '')
            authorization = authorizations.get(authorization_key)
            if not authorization:
                continue
            response_key = key(row.get('Response_ID'), f'RESP_DEMO_{index:04d}')
            payer_id = key(row.get('Payer_ID'), next(iter(payers)))
            response_id = uuid.uuid5(uuid.NAMESPACE_URL, f'pa360-response:{response_key}')
            PayerResponse.objects.update_or_create(
                response_id=response_id,
                defaults={
                    'authorization_request': authorization,
                    'payer': payers.get(payer_id) or next(iter(payers.values())),
                    'payer_name': clean(row.get('Payer_Name')) or 'Demo Payer',
                    'response_type': RESPONSE_MAP.get(clean(row.get('Response_Type')).upper(), PayerResponse.ResponseType.PENDING),
                    'response_date': datetime.combine(parse_date(row.get('Response_Date'), date.today()), datetime.min.time()),
                    'response_text': clean(row.get('Response_Text')) or 'Demo payer response',
                    'raw_response_data': {'source': 'workbook', 'response_id': response_key},
                },
            )

        self.stdout.write(self.style.SUCCESS(
            'Loaded demo workbook: '
            f'{Patient.objects.count()} patients, {Provider.objects.count()} providers, '
            f'{Payer.objects.count()} payers, {Plan.objects.count()} plans, '
            f'{Service.objects.count()} services, {AuthorizationRequest.objects.count()} authorizations, '
            f'{AuthorizationDocument.objects.count()} documents, {PayerResponse.objects.count()} responses.'
        ))
